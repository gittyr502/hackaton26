"""Excel export: RTL sheet, one row per form, layout from export.yaml."""
import csv
import io
import json
from typing import Any, Dict, List, Optional

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from pipeline.normalizers import to_display_date


def format_cell(field: str, value: Any, cfg: Dict[str, Any]) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return cfg.get("multi_value_separator", ", ").join(str(v) for v in value)
    if field == "receipt_date":
        return to_display_date(str(value), cfg.get("date_format", "DD/MM/YYYY")) or ""
    return str(value)


SOURCE_HEADER = "תמונת מקור"


def build_workbook(
    rows: List[Dict[str, Any]], cfg: Dict[str, Any], details: Optional[List[Dict[str, Any]]] = None,
    source_links: bool = False,
) -> Workbook:
    """rows: one dict per form, field name -> value (str | list | None).

    source_links: add a last column linking each row to its scanned page (row["_image_url"]).
    """
    wb = Workbook()
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)  # same data -> same bytes
    ws = wb.active
    ws.title = cfg.get("sheet_name", "Sheet1")[:31]
    ws.sheet_view.rightToLeft = bool(cfg.get("right_to_left", True))
    columns = cfg["columns"]

    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=col["header"])
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = 22
    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=format_cell(col["field"], row.get(col["field"]), cfg))
            cell.number_format = "@"  # everything as text: keeps ID leading zeros, no date coercion
    if source_links:
        c = len(columns) + 1
        head = ws.cell(row=1, column=c, value=SOURCE_HEADER)
        head.font = Font(bold=True)
        ws.column_dimensions[head.column_letter].width = 22
        for r, row in enumerate(rows, start=2):
            url = row.get("_image_url")
            if url:
                cell = ws.cell(row=r, column=c, value="פתח תמונה")
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    if details is not None:
        _add_details_sheet(wb, details, cfg)
    return wb


LEVEL_HE = {"ok": "תקין", "review": "לבדיקה", "empty": "ריק"}


def _add_details_sheet(wb: Workbook, details: List[Dict[str, Any]], cfg: Dict[str, Any]) -> None:
    """Quality-control sheet: per field level and model confidence, one row per form."""
    ws = wb.create_sheet("רמות ביטחון")
    ws.sheet_view.rightToLeft = bool(cfg.get("right_to_left", True))
    columns = cfg["columns"]
    headers = (["קובץ", "עמוד"] + [f"{c['header']} (רמה)" for c in columns]
               + [f"{c['header']} (ביטחון)" for c in columns])
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = 20
    for r, d in enumerate(details, start=2):
        ws.cell(row=r, column=1, value=d["filename"])
        ws.cell(row=r, column=2, value=d["page"])
        for i, col in enumerate(columns):
            f = d["fields"].get(col["field"])
            if f:
                ws.cell(row=r, column=3 + i, value=LEVEL_HE.get(f["level"], f["level"]))
                ws.cell(row=r, column=3 + len(columns) + i, value=f"{round(f['confidence'] * 100)}%")
    ws.freeze_panes = "A2"


def workbook_bytes(
    rows: List[Dict[str, Any]], cfg: Dict[str, Any], details: Optional[List[Dict[str, Any]]] = None,
    source_links: bool = False,
) -> bytes:
    buf = io.BytesIO()
    build_workbook(rows, cfg, details, source_links).save(buf)
    return buf.getvalue()


def encrypt_xlsx(data: bytes, password: str) -> bytes:
    """Open-with-password protection (Office agile encryption, AES-256)."""
    from msoffcrypto.format.ooxml import OOXMLFile

    out = io.BytesIO()
    OOXMLFile(io.BytesIO(data)).encrypt(password, out)
    return out.getvalue()


def csv_bytes(rows: List[Dict[str, Any]], cfg: Dict[str, Any]) -> bytes:
    """UTF-8 with BOM so Excel opens Hebrew correctly. Same columns/formatting as the xlsx."""
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    columns = cfg["columns"]
    w.writerow([c["header"] for c in columns])
    for row in rows:
        w.writerow([format_cell(c["field"], row.get(c["field"]), cfg) for c in columns])
    return buf.getvalue().encode("utf-8-sig")


def json_bytes(rows: List[Dict[str, Any]], cfg: Dict[str, Any]) -> bytes:
    """Rows keyed by field name; multi-value fields stay lists, dates formatted like the xlsx."""
    out = []
    for row in rows:
        item: Dict[str, Any] = {}
        for c in cfg["columns"]:
            v = row.get(c["field"])
            item[c["field"]] = v if isinstance(v, list) else (format_cell(c["field"], v, cfg) or None)
        out.append(item)
    return json.dumps(out, ensure_ascii=False, indent=2).encode("utf-8")
