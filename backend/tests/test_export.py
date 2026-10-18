import io

from openpyxl import load_workbook

import json

from export.excel import csv_bytes, json_bytes, workbook_bytes

CFG = {
    "sheet_name": "התחייבויות",
    "right_to_left": True,
    "date_format": "DD/MM/YYYY",
    "multi_value_separator": ", ",
    "columns": [
        {"field": "full_name", "header": "שם מלא"},
        {"field": "id_number", "header": "ת.ז."},
        {"field": "receipt_date", "header": "תאריך קבלה"},
        {"field": "commitment_numbers", "header": "מספרי התחייבות"},
    ],
}


def test_export_layout_and_formats():
    rows = [
        {"full_name": "דנה", "id_number": "18", "receipt_date": "2026-03-05",
         "commitment_numbers": ["111", "222"]},
        {"full_name": None, "id_number": None, "receipt_date": None, "commitment_numbers": None},
    ]
    wb = load_workbook(io.BytesIO(workbook_bytes(rows, CFG)))
    ws = wb.active
    assert ws.title == "התחייבויות"
    assert ws.sheet_view.rightToLeft
    assert [c.value for c in ws[1]] == ["שם מלא", "ת.ז.", "תאריך קבלה", "מספרי התחייבות"]
    assert ws["B2"].value == "18"
    assert ws["B2"].data_type == "s"
    assert ws["B2"].number_format == "@"
    assert ws["C2"].value == "05/03/2026"
    assert ws["D2"].value == "111, 222"
    assert ws["A3"].value in (None, "")


ROWS = [{"full_name": "דנה", "id_number": "18", "receipt_date": "2026-03-05",
         "commitment_numbers": ["111", "222"]}]


def test_csv_export():
    lines = csv_bytes(ROWS, CFG).decode("utf-8-sig").strip().splitlines()
    assert lines[0] == "שם מלא,ת.ז.,תאריך קבלה,מספרי התחייבות"
    assert lines[1] == 'דנה,18,05/03/2026,"111, 222"'


def test_json_export():
    data = json.loads(json_bytes(ROWS, CFG))
    assert data[0]["id_number"] == "18"
    assert data[0]["commitment_numbers"] == ["111", "222"]


def test_details_sheet():
    details = [{"filename": "a.pdf", "page": 1,
                "fields": {"full_name": {"level": "review", "confidence": 0.62, "manually_corrected": False}}}]
    ws = load_workbook(io.BytesIO(workbook_bytes(ROWS, CFG, details)))["רמות ביטחון"]
    assert ws["A2"].value == "a.pdf"
    assert ws["C2"].value == "לבדיקה"
    assert ws["G2"].value == "62%"
