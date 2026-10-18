"""API: corrections persist, export summary counts review fields, delete removes everything."""
import json

from fastapi.testclient import TestClient
from sqlmodel import Session

from app import db
from app.main import app
from app.models import Batch, Field, Form, SourceFile


def _seed(settings, workflow=None):
    batch_dir = settings.batches_dir / "1"
    batch_dir.mkdir(parents=True)
    with Session(db.get_engine()) as s:
        b = Batch(status="done")
        s.add(b)
        s.commit()
        sf = SourceFile(batch_id=b.id, filename="a.png", path=str(batch_dir / "a.png"), status="done")
        s.add(sf)
        s.commit()
        fm = Form(batch_id=b.id, source_file_id=sf.id, status="done", workflow=workflow)
        s.add(fm)
        s.commit()
        fld = Field(form_id=fm.id, name="receipt_date", value=json.dumps("2026-13-01"), level="review",
                    reasons='[{"kind": "validation", "code": "invalid_date"}]')
        s.add(fld)
        s.commit()
        return b.id, fld.id


def test_correction_persists_and_export(tmp_env):
    batch_id, field_id = _seed(tmp_env)
    with TestClient(app) as client:
        assert client.get(f"/api/batches/{batch_id}/export/summary").json()["review"] == 1

        r = client.patch(f"/api/fields/{field_id}", json={"value": "01/03/2026"})
        assert r.status_code == 200
        assert r.json()["value"] == "2026-03-01" and r.json()["level"] == "ok"

        detail = client.get(f"/api/batches/{batch_id}").json()  # "after refresh"
        fld = detail["forms"][0]["fields"]["receipt_date"]
        assert fld["value"] == "2026-03-01" and fld["manually_corrected"]
        assert client.get(f"/api/batches/{batch_id}/export/summary").json()["review"] == 0

        x = client.get(f"/api/batches/{batch_id}/export")
        assert x.status_code == 200 and x.content[:2] == b"PK"

        assert client.delete(f"/api/batches/{batch_id}").status_code == 200
        assert client.get(f"/api/batches/{batch_id}").status_code == 404
        assert not (tmp_env.batches_dir / str(batch_id)).exists()


def test_field_box_parsing():
    from pipeline.stages.extract import parse_field

    ok = parse_field({"value": "123", "confidence": 0.9, "legible": True, "box": [0.1, 0.2, 0.4, 0.25]})
    assert ok.box == [0.1, 0.2, 0.4, 0.25]
    assert parse_field({"value": "1", "confidence": 1, "box": [0.5, 0.5, 0.5, 0.5]}).box is None
    # an empty field keeps its box so it can be re-read from a zoomed crop
    assert parse_field({"value": None, "confidence": 0, "box": [0, 0, 1, 1]}).box == [0, 0, 1, 1]
    assert parse_field({"value": "1", "confidence": 1, "box": "oops"}).box is None


def test_export_options_saved_in_db_and_protected_download(tmp_env):
    import io

    import msoffcrypto
    from openpyxl import load_workbook

    batch_id, _ = _seed(tmp_env, workflow="approved")  # only approved forms are exported
    with TestClient(app) as client:
        s = client.get(f"/api/batches/{batch_id}/export/summary").json()
        assert s["scan_number"].startswith("MED-") and s["file"]["filename"].endswith(".xlsx")
        assert s["options"]["include_details"] and s["file"]["sheets"] == 2

        r = client.put(f"/api/batches/{batch_id}/export/options",
                       json={"format": "xlsx", "include_details": False,
                             "include_source_links": True, "password_protect": True})
        assert r.json()["file"]["sheets"] == 1
        assert client.get(f"/api/batches/{batch_id}/export/summary").json()["options"]["password_protect"]

        assert client.post(f"/api/batches/{batch_id}/export", json={}).status_code == 400  # no password
        x = client.post(f"/api/batches/{batch_id}/export", json={"password": "s3cret", "link_base": "http://x"})
        assert x.status_code == 200
        out = io.BytesIO()
        f = msoffcrypto.OfficeFile(io.BytesIO(x.content))
        f.load_key(password="s3cret")
        f.decrypt(out)
        ws = load_workbook(out).active
        assert ws.cell(1, 7).value == "תמונת מקור" and ws.cell(2, 7).hyperlink.target.startswith("http://x/api/forms/")
        assert client.get(f"/api/batches/{batch_id}/export/summary").json()["last_export"]["protected"]


def test_form_image_and_config_come_from_db(tmp_env):
    batch_id, _ = _seed(tmp_env)
    with Session(db.get_engine()) as s:
        fm = s.get(Form, 1)
        fm.image_data = b"\x89PNG-from-db"
        s.add(fm)
        s.commit()
    with TestClient(app) as client:
        assert client.get("/api/forms/1/image").content == b"\x89PNG-from-db"
        cfg = client.get("/api/config/export").json()
        cfg["columns"][0]["header"] = "פרטי"
        client.put("/api/config/export", json=cfg)
        assert client.get(f"/api/batches/{batch_id}/export/summary").json()["columns"][0]["header"] == "פרטי"
